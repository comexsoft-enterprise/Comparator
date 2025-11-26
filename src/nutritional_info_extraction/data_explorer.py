import pandas as pd
import numpy as np
import logging
from colorlog import ColoredFormatter

# Formato condicional con openpyxl
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

# Configuración de logs con colores
formatter = ColoredFormatter(
    '%(log_color)s%(asctime)s - %(levelname)s - %(message)s',
    datefmt='%H:%M:%S',
    log_colors={
        'DEBUG':    'cyan',
        'INFO':     'green',
        'WARNING':  'yellow',
        'ERROR':    'red',
        'CRITICAL': 'bold_red',
    }
)
handler = logging.StreamHandler()
handler.setFormatter(formatter)
logger = logging.getLogger('nutri_logger')
logger.setLevel(logging.INFO)
logger.handlers = []
logger.addHandler(handler)

def log_section(title):
    logger.info('\n' + '='*60)
    logger.info(f' {title} ')
    logger.info('='*60)

log_section('Extracción y análisis de información nutricional')

# Rutas de los archivos
eroski_path = "src/nutritional_info_extraction/eroski_01013_translated_1.csv"
makro_path = "src/nutritional_info_extraction/makro_01013_translated_1.csv"

# Cargar datos
log_section('Cargando archivos CSV')
eroski = pd.read_csv(eroski_path, sep=";")
makro = pd.read_csv(makro_path, sep=";")
logger.info(f'Eroski shape: {eroski.shape}')
logger.info(f'Makro shape: {makro.shape}')

log_section('Primeras filas de Eroski')
logger.info(f"\n{eroski.head(3)}")
log_section('Primeras filas de Makro')
logger.info(f"\n{makro.head(3)}")

# Identificar columnas nutricionales
def get_nutrition_columns(df):
    return [col for col in df.columns if "nutri" in col.lower()]

nutritional_columns_eroski = get_nutrition_columns(eroski)
nutritional_columns_makro = get_nutrition_columns(makro)
log_section('Columnas nutricionales detectadas')
logger.info(f'Nutritional columns in eroski: {nutritional_columns_eroski} ({len(nutritional_columns_eroski)})')
logger.info(f'Nutritional columns in makro: {nutritional_columns_makro} ({len(nutritional_columns_makro)})')

# Analizar columnas nutricionales
def analyze_nutrition_columns(df, name):
    nutrition_cols = get_nutrition_columns(df)
    log_section(f'Análisis de columnas nutricionales: {name}')
    for col in nutrition_cols:
        col_type = df[col].dtype
        unique_vals = df[col].dropna().unique()
        sample_vals = unique_vals[:5]
        n_nulos = df[col].isnull().sum()
        logger.info(f"Columna: {col}")
        logger.info(f"  Tipo: {col_type}")
        logger.info(f"  Ejemplo valores: {sample_vals}")
        logger.info(f"  Nulos: {n_nulos} / {len(df)}")

analyze_nutrition_columns(eroski, "Eroski")
analyze_nutrition_columns(makro, "Makro")

# Limpieza y exportación

def clean_nutrition_df(df, source_name):
    nutrition_cols = get_nutrition_columns(df)
    # Añadir uuid y product_name si existen
    extra_cols = []
    for col in ["uuid", "product_name"]:
        if col in df.columns:
            extra_cols.append(col)
    nutrition_df = df[nutrition_cols + extra_cols].copy()
    nutrition_df = nutrition_df.dropna(axis=1, how='all')
    nutrition_df = nutrition_df.fillna("")
    nutrition_df["source"] = source_name
    return nutrition_df

nutrition_eroski = clean_nutrition_df(eroski, "eroski")
nutrition_makro = clean_nutrition_df(makro, "makro")
nutrition_all = pd.concat([nutrition_eroski, nutrition_makro], ignore_index=True)

# Unificar unidades nutricionales
unit_standard = {
    'nutrition_information_calories': ('kcal', {'kcal': 1, 'kj': 0.239006}),
    'nutrition_information_fat': ('g', {'g': 1}),
    'nutrition_information_carbohydrates': ('g', {'g': 1}),
    'nutrition_information_fiber': ('g', {'g': 1}),
    'nutrition_information_protein': ('g', {'g': 1}),
    'nutrition_information_salt': ('g', {'g': 1}),
    'nutrition_information_sugars': ('g', {'g': 1}),
    'nutrition_information_saturatedfattyacids': ('g', {'g': 1}),
    'nutrition_information_monounsaturatedfattyacids': ('g', {'g': 1}),
    'nutrition_information_polyunsaturatedgradeacids': ('g', {'g': 1}),
}

def unify_units(df):
    for base, (std_unit, factors) in unit_standard.items():
        value_col = f'{base}_value'
        unit_col = f'{base}_unit'
        if value_col in df.columns and unit_col in df.columns:
            def convert(row):
                val = row[value_col]
                unit = row[unit_col]
                # Estandarizar: convertir comas a puntos y a float
                if isinstance(val, str):
                    val = val.replace(',', '.')
                    try:
                        val = float(val)
                    except Exception:
                        pass
                if pd.isnull(val) or pd.isnull(unit) or unit == std_unit:
                    return val
                try:
                    return float(val) * factors.get(unit.lower(), 1)
                except Exception:
                    return val
            df[value_col] = df.apply(convert, axis=1)
            df[unit_col] = std_unit
    return df

nutrition_all = unify_units(nutrition_all)
output_path = "src/nutritional_info_extraction/nutrition_info_combined.xlsx"
# Ordenar columnas para coherencia: primero valores, luego unidades, luego source
def order_nutrition_columns(df):
    value_cols = [col for col in df.columns if col.endswith('_value')]
    unit_cols = [col for col in df.columns if col.endswith('_unit')]
    nutrients = sorted(set([col.replace('nutrition_information_', '').replace('_value', '').replace('_unit', '') for col in value_cols + unit_cols]))
    ordered_cols = []
    for nutrient in nutrients:
        value_col = f'nutrition_information_{nutrient}_value'
        unit_col = f'nutrition_information_{nutrient}_unit'
        if value_col in df.columns:
            ordered_cols.append(value_col)
        if unit_col in df.columns:
            ordered_cols.append(unit_col)
    # Añadir las columnas que no son value/unit ni source
    other_cols = [col for col in df.columns if col not in ordered_cols and col != 'source']
    ordered_cols += other_cols
    # Añadir 'source' al final
    if 'source' in df.columns:
        ordered_cols.append('source')
    return df[ordered_cols]

# Aplicar orden antes de exportar
nutrition_all_ordered = order_nutrition_columns(nutrition_all)
nutrition_all_ordered.to_excel(output_path, index=False)
nutrition_all_ordered.to_csv(output_path.replace('.xlsx', '.csv'), index=False, sep=";")
wb = load_workbook(output_path)
ws = wb.active

# Buscar la columna 'source'
source_col_idx = None
for idx, cell in enumerate(ws[1], 1):
    if cell.value == 'source':
        source_col_idx = idx
        break

eroski_fill = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')  # Verde claro
makro_fill = PatternFill(start_color='FFEB9C', end_color='FFEB9C', fill_type='solid')  # Amarillo claro

for row in ws.iter_rows(min_row=2, min_col=1, max_col=ws.max_column):
    source = row[source_col_idx-1].value
    if source == 'eroski':
        for cell in row:
            cell.fill = eroski_fill
    elif source == 'makro':
        for cell in row:
            cell.fill = makro_fill

wb.save(output_path)
log_section('Exportación a Excel con formato de colores por supermercado')
logger.info(f"Exportado a {output_path} con colores diferenciados para eroski y makro")

# Análisis del Excel final

final_df = pd.read_excel(output_path)
# Considerar como incompleto: '', None, NaN
def completion_percent(col):
    return (final_df[col].replace('', np.nan).notnull().mean() * 100)
completitud = {col: completion_percent(col) for col in final_df.columns}
print("\nResumen de completitud por columna:")
for col, pct in completitud.items():
    print(f"{col}: {pct:.2f}%")
print("\nFilas por origen:")
for origen, count in final_df['source'].value_counts().items():
    print(f"{origen}: {count}")